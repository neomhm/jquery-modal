"""
gen/realfile.py - write a generated sheet as a REAL .xlsx or .csv file and
read it back with sheets.load() (section 10.1: 1% of tasks). The preview
of the file must be identical to the preview of the generated sheet.

    write_xlsx(sheet, path)             typed values and number formats
    write_csv(sheet, path, delimiter)   all text, named after its sheet
    round_trip(task, folder) -> (ok, detail)
"""
import csv
import datetime
import pathlib

import sheets

BAD_NAME = '[]:*?/\\'


def write_xlsx(sheet, path):
    import openpyxl
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = sheet.name[:31]
    for i, row in enumerate(sheet.rows):
        for j, v in enumerate(row):
            if v is None:
                continue
            cell = ws.cell(row=i + 1, column=j + 1, value=v)
            fmt = None
            if sheet.formats and i < len(sheet.formats) and \
                    sheet.formats[i] and j < len(sheet.formats[i]):
                fmt = sheet.formats[i][j]
            if isinstance(v, datetime.timedelta) and not fmt:
                fmt = "[h]:mm"
            if fmt:
                cell.number_format = fmt
    book.save(path)


def write_csv(sheet, path, delimiter=",", encoding="utf-8"):
    """Every row padded to the sheet's width, as Excel writes CSV (a
    title row is 'Tarifs 2026;;;;'), so the delimiter can be sniffed."""
    width = max([len(r) for r in sheet.rows] + [1])
    with open(path, "w", encoding=encoding, newline="") as f:
        w = csv.writer(f, delimiter=delimiter, quoting=csv.QUOTE_MINIMAL)
        for row in sheet.rows:
            cells = ["" if v is None else str(v) for v in row]
            w.writerow(cells + [""] * (width - len(cells)))


def write_xls(sheet, path):
    """An old Excel .xls file (xlwt, vendor/): typed values and number
    formats, as write_xlsx."""
    import xlwt
    book = xlwt.Workbook(encoding="utf-8")
    ws = book.add_sheet("".join("_" if ch in BAD_NAME else ch
                                for ch in sheet.name)[:31] or "Sheet1")
    styles = {}

    def style(fmt):
        if fmt not in styles:
            styles[fmt] = xlwt.easyxf(num_format_str=fmt)
        return styles[fmt]
    for i, row in enumerate(sheet.rows):
        for j, v in enumerate(row):
            if v is None:
                continue
            fmt = None
            if sheet.formats and i < len(sheet.formats) and \
                    sheet.formats[i] and j < len(sheet.formats[i]):
                fmt = sheet.formats[i][j]
            if isinstance(v, datetime.timedelta):
                v, fmt = v.total_seconds() / 86400.0, fmt or "[h]:mm"
            elif isinstance(v, datetime.datetime) and not fmt:
                fmt = "yyyy-mm-dd hh:mm"
            elif isinstance(v, datetime.date) and not fmt:
                fmt = "yyyy-mm-dd"
            elif isinstance(v, datetime.time) and not fmt:
                fmt = "hh:mm"
            if fmt and fmt != "General":
                ws.write(i, j, v, style(fmt))
            else:
                ws.write(i, j, v)
    book.save(str(path))


def _ods_cell(v, fmt, locale=None):
    """One table:table-cell of an .ods file."""
    from xml.sax.saxutils import escape, quoteattr
    if v is None:
        return "<table:table-cell/>"
    text = escape(str(v))
    if isinstance(v, bool):
        return ('<table:table-cell office:value-type="boolean" '
                'office:boolean-value="%s"><text:p>%s</text:p>'
                '</table:table-cell>' % ("true" if v else "false",
                                         "TRUE" if v else "FALSE"))
    if isinstance(v, (int, float)):
        f = fmt or ""
        if "%" in f:
            kind = 'office:value-type="percentage"'
        else:
            code = _currency_of_format(f, locale)
            kind = ('office:value-type="currency" office:currency=%s'
                    % quoteattr(code)) if code else \
                'office:value-type="float"'
        return ('<table:table-cell %s office:value="%r"><text:p>%s</text:p>'
                '</table:table-cell>' % (kind, v, text))
    if isinstance(v, datetime.datetime):
        return ('<table:table-cell office:value-type="date" '
                'office:date-value="%s"><text:p>%s</text:p>'
                '</table:table-cell>' % (
                    v.isoformat() if v.time() != datetime.time(0)
                    else v.date().isoformat(), text))
    if isinstance(v, datetime.date):
        return ('<table:table-cell office:value-type="date" '
                'office:date-value="%s"><text:p>%s</text:p>'
                '</table:table-cell>' % (v.isoformat(), text))
    if isinstance(v, datetime.time):
        return ('<table:table-cell office:value-type="time" '
                'office:time-value="PT%02dH%02dM%02dS"><text:p>%s</text:p>'
                '</table:table-cell>' % (v.hour, v.minute, v.second, text))
    if isinstance(v, datetime.timedelta):
        total = int(v.total_seconds())
        return ('<table:table-cell office:value-type="time" '
                'office:time-value="PT%dH%02dM%02dS"><text:p>%s</text:p>'
                '</table:table-cell>' % (total // 3600, total // 60 % 60,
                                         total % 60, text))
    lines = str(v).split("\n")
    return ('<table:table-cell office:value-type="string">%s'
            '</table:table-cell>' % "".join(
                "<text:p>%s</text:p>" % escape(x) for x in lines))


def _currency_of_format(fmt, locale=None):
    """The ISO code an Excel number format shows ('#,##0.00 [$€-40C]'),
    read in the sheet's locale ("¥" is CNY in China, JPY in Japan), or
    None."""
    if not fmt:
        return None
    import helpers as H
    ok, code = H.currency(fmt, locale)
    return code if ok else None


def write_ods(sheet, path, locale=None):
    """A LibreOffice .ods file (a zip of XML, standard library): typed
    values; a number with a currency or percent format becomes a currency
    or percentage cell (ODS has no Excel format strings)."""
    import zipfile
    from xml.sax.saxutils import quoteattr
    ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
          'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"')
    rows = []
    for i, row in enumerate(sheet.rows):
        cells = []
        for j, v in enumerate(row):
            fmt = None
            if sheet.formats and i < len(sheet.formats) and \
                    sheet.formats[i] and j < len(sheet.formats[i]):
                fmt = sheet.formats[i][j]
            cells.append(_ods_cell(v, fmt, locale))
        rows.append("<table:table-row>%s</table:table-row>" % (
            "".join(cells) or "<table:table-cell/>"))
    content = ('<?xml version="1.0" encoding="UTF-8"?>'
               '<office:document-content %s office:version="1.2">'
               '<office:body><office:spreadsheet><table:table '
               'table:name=%s>%s</table:table></office:spreadsheet>'
               '</office:body></office:document-content>' % (
                   ns, quoteattr(sheet.name), "".join(rows)))
    manifest = ('<?xml version="1.0" encoding="UTF-8"?>'
                '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:'
                'opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
                '<manifest:file-entry manifest:full-path="/" '
                'manifest:media-type="application/vnd.oasis.opendocument.'
                'spreadsheet"/><manifest:file-entry manifest:full-path='
                '"content.xml" manifest:media-type="text/xml"/>'
                '</manifest:manifest>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"),
                   "application/vnd.oasis.opendocument.spreadsheet",
                   compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/manifest.xml", manifest,
                   compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)


def write_pdf(sheet, path, right=(), **options):
    """A PDF table of the sheet's cells as texts (gen/pdfwrite.py)."""
    from gen import pdfwrite
    pdfwrite.write_pdf(path, [["" if v is None else str(v) for v in r]
                              for r in sheet.rows], right=right, **options)


def round_trip(task, folder):
    """-> (True, '') when the file's preview equals the task's preview."""
    from gen import data as D
    from gen.tasks import sheet_from_task
    sheet = sheet_from_task(task)
    if task.get("answer") == "multi":
        return True, ""            # several tables: checked when drawn
    folder = pathlib.Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    name = "".join("_" if ch in BAD_NAME else ch for ch in sheet.name)
    if task["format"] == "xlsx":
        path = folder / ("%s.xlsx" % task["id"])
        write_xlsx(sheet, path)
    elif task["format"] in ("xls", "ods", "pdf"):
        sub = folder / task["id"]
        sub.mkdir(exist_ok=True)
        path = sub / ("%s.%s" % (name.strip() or "Sheet1", task["format"]))
        if task["format"] == "xls":
            write_xls(sheet, path)
        elif task["format"] == "ods":
            write_ods(sheet, path, task["locale"])
        else:
            from gen import filetypes
            write_pdf(sheet, path, right=filetypes._numeric_columns(sheet))
    else:
        sub = folder / task["id"]
        sub.mkdir(exist_ok=True)
        path = sub / ("%s.csv" % name)        # a CSV is named after its sheet
        loc = D.locale(task["locale"])
        write_csv(sheet, path, loc.get("csv_delimiter") or ",",
                  "utf-8-sig" if task["lang"] in ("ja", "zh", "ko", "ar",
                                                  "hi", "ru") else "utf-8")
    back = sheets.load(path, task["locale"])
    if not back:
        return False, "not read back"
    back[0].name = sheet.name          # a PDF is named after its file
    small, _ = sheets.compact(back[0])
    again = sheets.preview(small, task["targets"], task["locale"])
    if again != task["preview"]:
        a, b = again.splitlines(), task["preview"].splitlines()
        for x, y in zip(a, b):
            if x != y:
                return False, "file: %s | task: %s" % (x[:120], y[:120])
        return False, "different length"
    return True, ""
